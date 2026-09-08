from pathlib import Path

import pytest
from conftest import CharacterTokenizer

from backend.app.chunking import split_chunks
from backend.app.parsers import Block, parse_document


def test_txt_utf8_and_gb18030(tmp_path, settings):
    path = tmp_path / "policy.txt"
    for encoding in ("utf-8-sig", "gb18030"):
        path.write_bytes("北京住宿标准 600 元。\n\n餐补 100 元。".encode(encoding))
        blocks = parse_document(path, settings)
        assert len(blocks) == 2
        assert "600" in blocks[0].text


def test_markdown_headings_preserved(settings):
    blocks = parse_document(Path("samples/公司测试制度.md"), settings)
    assert any(b.section == "4.3 酒店住宿标准" and "600" in b.text for b in blocks)


def test_chunk_boundaries_offsets_and_coverage():
    tokenizer = CharacterTokenizer()
    text = "北京住宿标准为每晚600元。" * 30
    blocks = [Block(text, 3, "住宿"), Block("餐补100元", 4, "餐饮")]
    chunks = split_chunks(blocks, "doc-id", "政策.pdf", tokenizer, 80, 12)
    assert all(len(c.text) <= 80 for c in chunks)
    assert all(c.page == 3 and c.section == "住宿" for c in chunks[:-1])
    assert chunks[-1].page == 4 and chunks[-1].section == "餐饮"
    reconstructed = chunks[0].text + "".join(c.text[12:] for c in chunks[1:-1])
    assert reconstructed == text
    assert len({c.chunk_id for c in chunks}) == len(chunks)
    assert [c.chunk_index for c in chunks] == list(range(len(chunks)))
    assert chunks[0].chunk_id == split_chunks(blocks, "doc-id", "政策.pdf", tokenizer, 80, 12)[0].chunk_id


def test_docx_paragraph_table_order(tmp_path, settings):
    from docx import Document

    doc = Document()
    doc.add_heading("住宿标准", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "北京"
    table.cell(0, 1).text = "600元"
    doc.add_paragraph("后续说明")
    path = tmp_path / "policy.docx"
    doc.save(path)
    blocks = parse_document(path, settings)
    assert [b.text for b in blocks] == ["住宿标准", "北京 | 600元", "后续说明"]
    assert all(b.section == "住宿标准" and b.page is None for b in blocks)


def test_pdf_page_numbers_and_scan_error(tmp_path, settings):
    import pymupdf

    path = tmp_path / "policy.pdf"
    with pymupdf.open() as doc:
        for text in ("Beijing 600", "Shanghai 550"):
            doc.new_page().insert_text((72, 72), text)
        doc.save(path)
    blocks = parse_document(path, settings)
    assert [(b.page, b.text) for b in blocks] == [(1, "Beijing 600"), (2, "Shanghai 550")]
    with pymupdf.open() as doc:
        doc.new_page()
        doc.save(tmp_path / "scan.pdf")
    with pytest.raises(ValueError, match="OCR"):
        parse_document(tmp_path / "scan.pdf", settings)


def test_empty_and_oversized_parsing(tmp_path, settings):
    path = tmp_path / "empty.txt"
    path.write_text("   ")
    with pytest.raises(ValueError, match="未提取"):
        parse_document(path, settings)
    path.write_text("a" * 100)
    settings.max_parsed_chars = 50
    with pytest.raises(ValueError, match="过大"):
        parse_document(path, settings)
